"""Regressões dos cruzamentos GIR/GGR/GOP sem documentos pessoais no repositório."""
from decimal import Decimal as D
from pathlib import Path
import os
import unittest

from conciliacao import conciliar, ler_dcb
from ponte import Arquivo, processar
from sistemas import ler_sistema, cruzar_sistemas


def pagamento(boleto='123456', valor='10.00', sistema='GGR'):
    return dict(boleto=boleto, valor=D(valor), sistema=sistema, nome='Cliente teste',
                documento='1/2026', contrato='', data_pagamento='28/09/2026', pagina=3)


def titulo(boleto='123456', valor='10.00', data='28/09/2026', carteira='6'):
    return dict(nosso_numero=carteira+boleto+'07000', valor=D(valor), data=data,
                data_ocorrencia=data, nome='Cliente teste')


def relatorio(*registros, sistema='GGR'):
    return dict(sistema=sistema, registros=list(registros))


class CruzamentosTest(unittest.TestCase):
    def test_soma_parcelas_e_consolida_fora_do_gai(self):
        r = relatorio(pagamento(valor='4.00'), pagamento(valor='6.00'))
        dcb = {'liquidados': [titulo()]}
        bb = {'registros': [titulo()]}
        resultado = cruzar_sistemas([r], dcb, bb)
        self.assertEqual(len(r['linhas']), 1)
        self.assertEqual(r['linhas'][0]['situacao_dcb'], 'Conferido')
        self.assertEqual(resultado['quantidade_pendentes'], 0)
        self.assertEqual(resultado['linhas'][0]['sistema'], 'GGR')

    def test_data_numero_e_valor_sao_conferidos(self):
        r = relatorio(pagamento())
        cruzar_sistemas([r], {'liquidados': [titulo(data='29/09/2026'), titulo(boleto='999999')]})
        self.assertEqual(r['linhas'][0]['situacao_dcb'], 'Não encontrado')
        cruzar_sistemas([r], {'liquidados': [titulo(valor='11.00')]})
        self.assertEqual(r['linhas'][0]['situacao_dcb'], 'Valor divergente')
        self.assertEqual(r['linhas'][0]['situacao_bolebarra'], 'Não enviado')

    def test_carteiras_e_sistemas_ambiguos_nao_conferem(self):
        r = relatorio(pagamento())
        cruzar_sistemas([r], {'liquidados': [titulo(valor='5'), titulo(valor='5', carteira='1')]})
        self.assertEqual(r['linhas'][0]['situacao_dcb'], 'Identificação ambígua')
        outro = relatorio(pagamento(sistema='GOP'), sistema='GOP')
        resultado = cruzar_sistemas([r, outro], bolebarras={'registros': [titulo(valor='20')]})
        self.assertEqual(resultado['linhas'][0]['situacao'], 'Identificação ambígua')

    def test_gai_tambem_entra_no_consolidado(self):
        r = relatorio(pagamento())
        bb = {'registros': [titulo(), titulo(boleto='654321')]}
        gai = {'parcelas': [pagamento(boleto='654321', sistema='GAI')]}
        resultado = cruzar_sistemas([r], bolebarras=bb, gai_lidas=gai)
        self.assertEqual(resultado['quantidade_pendentes'], 0)
        self.assertEqual(resultado['linhas'][1]['sistema'], 'GAI')


PASTA = Path(os.environ.get('CONCILIACAO_TESTE3', 'tmp/fixtures/TESTE 3'))

@unittest.skipUnless(PASTA.is_dir(), 'Defina CONCILIACAO_TESTE3 para testar os PDFs reais')
class PDFsReaisTest(unittest.TestCase):
    def test_tabelas_destacadas_totais_e_dcb(self):
        esperado = {'GIR': (6, D('74517.16'), [2]),
                    'GGR': (7, D('45140.74'), [3]),
                    'GOP': (1, D('116043.00'), [3])}
        dcb = ler_dcb(next(PASTA.glob('DCB*')))
        for sistema, (qtd, total, paginas) in esperado.items():
            with self.subTest(sistema=sistema):
                arquivo = next(PASTA.glob('*'+sistema+'*.pdf'))
                # Mesma entrada de bytes usada no navegador.
                r = ler_sistema(Arquivo('relatorio.pdf', arquivo.read_bytes()), sistema)
                self.assertEqual((r['quantidade'], r['total'], r['paginas']), (qtd, total, paginas))
                self.assertEqual(r['total'], r['total_impresso'])
                self.assertEqual(r['avisos'], [])
                self.assertTrue(all(x['valor'] > 0 for x in r['registros']))
                cruzar_sistemas([r], dcb)
                self.assertTrue(all(x['situacao_dcb'] == 'Conferido' for x in r['linhas']))
        with self.assertRaises(ValueError):
            ler_sistema(next(PASTA.glob('*GIR*.pdf')), 'GGR')

    def test_entrada_navegador_e_excel(self):
        import json
        from io import BytesIO
        from openpyxl import load_workbook
        from ponte import exportar
        arquivos = {s.lower(): {'nome': 'relatorio.pdf', 'dados': next(PASTA.glob('*'+s+'*.pdf')).read_bytes()}
                    for s in ('GIR', 'GGR', 'GOP')}
        arquivos['dcb'] = {'nome': 'retorno.txt', 'dados': next(PASTA.glob('DCB*')).read_bytes()}
        resultado = json.loads(processar(arquivos))
        self.assertNotIn('erro', resultado)
        self.assertEqual(len(resultado['sistemas']), 3)
        livro = load_workbook(BytesIO(exportar()), data_only=True)
        for sistema, qtd in [('GIR', 6), ('GGR', 7), ('GOP', 1)]:
            aba = livro[sistema+' x banco']
            self.assertEqual(aba.max_row, qtd+1)
            self.assertTrue(all(aba.cell(i, 7).value > 0 for i in range(2, qtd+2)))
            self.assertTrue(all(aba.cell(i, 9).value == 'Conferido' for i in range(2, qtd+2)))


if __name__ == '__main__':
    unittest.main()
